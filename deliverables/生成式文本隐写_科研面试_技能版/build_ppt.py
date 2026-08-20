# -*- coding: utf-8 -*-
"""Build the 10-slide academic interview deck (16:9, native editable slides).

Layout system (user-confirmed combination):
- A 科研答辩 · 分栏证据版：左文右图分栏、严格网格、顶部结论条、底部来源/边界行
- D 时间线叙事主线：页眉带 6 段进度条（背景→问题→路线→方法→验证→总结）
- 结果页大数字卡；第 6 页逐步推导条

Image generation is deferred (九秋 backend pending); concept visuals are dashed
placeholders that auto-swap to PNGs placed in assets/ on rebuild. Real data
figures from the repo are embedded directly.
"""

import os
import re

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, "生成式文本隐写_科研面试_技能版.pptx")
SPEECH = os.path.join(ROOT, "speech.md")
ASSETS = os.path.join(ROOT, "assets")

REPO = r"C:\Users\41462\Documents\隐写"
FIG2 = os.path.join(REPO, "figures", "fig2_main_scale.png")
FIG4 = os.path.join(REPO, "deliverables", "assets", "fig4_precision_direction.png")
ARCH = os.path.join(REPO, "paper", "ccfa_bounded_decision", "figures", "figure_01_architecture.png")

# 科研答辩风 palette
BLUE = RGBColor(0x00, 0x3F, 0x8F)
BLUE2 = RGBColor(0x0B, 0x5C, 0xAD)
PALE = RGBColor(0xEA, 0xF2, 0xFF)
RED = RGBColor(0xB5, 0x12, 0x1B)
TEAL = RGBColor(0x0C, 0x8F, 0x88)
ORANGE = RGBColor(0xD8, 0x78, 0x1F)
BLACK = RGBColor(0x11, 0x11, 0x11)
GRAY = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x5A, 0x64, 0x72)
BORDER = RGBColor(0xD8, 0xDE, 0xE8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL_PALE = RGBColor(0xE6, 0xF4, 0xF2)
RED_PALE = RGBColor(0xFB, 0xEC, 0xEC)
ORANGE_PALE = RGBColor(0xFD, 0xF3, 0xE7)

STAGES = ["背景", "问题", "路线", "方法", "验证", "总结"]
FONT = "Microsoft YaHei"


def set_run(run, size, bold=False, color=BLACK, italic=False, font=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", font)


def add_text(slide, text, x, y, w, h, size=13, bold=False, color=BLACK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, italic=False,
             spacing=1.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        set_run(r, size, bold, color, italic)
    return box


def shape(slide, st, x, y, w, h, fill=WHITE, line=BORDER, width=0.8, dash=None, radius=None):
    sp = slide.shapes.add_shape(st, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(width)
        if dash is not None:
            ln = sp.line._get_or_add_ln()
            prst = ln.makeelement(qn("a:prstDash"), {"val": dash})
            ln.append(prst)
    sp.shadow.inherit = False
    if radius is not None and st == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def rect(slide, x, y, w, h, fill, line=None):
    return shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line)


def rrect(slide, x, y, w, h, fill=WHITE, line=BORDER, radius=0.06):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, radius=radius)


def chevron(slide, x, y, w, color=BLUE2):
    return shape(slide, MSO_SHAPE.CHEVRON, x, y, w, 0.26, color, color)


def card(slide, x, y, w, h, heading, body, accent=BLUE, fill=WHITE,
         hsize=12.5, bsize=9.5, head_color=None):
    rrect(slide, x, y, w, h, fill, BORDER)
    rect(slide, x, y, 0.055, h, accent)
    add_text(slide, heading, x + 0.2, y + 0.13, w - 0.38, 0.3, hsize, True,
             head_color or accent)
    add_text(slide, body, x + 0.2, y + 0.5, w - 0.38, h - 0.64, bsize, False,
             GRAY, anchor=MSO_ANCHOR.TOP, spacing=1.15)


def picture_fit(slide, path, x, y, max_w, max_h, border=False):
    from PIL import Image
    with Image.open(path) as im:
        w, h = im.size
    ratio = min(max_w / w, max_h / h)
    dw, dh = w * ratio, h * ratio
    px = x + (max_w - dw) / 2
    py = y + (max_h - dh) / 2
    if border:
        rrect(slide, x - 0.04, y - 0.04, max_w + 0.08, max_h + 0.08, WHITE, BORDER, radius=0.03)
    slide.shapes.add_picture(path, Inches(px), Inches(py), Inches(dw), Inches(dh))


def placeholder(slide, x, y, w, h, filenames, note="插图预留位（待生图后替换）"):
    for f in filenames:
        p = os.path.join(ASSETS, f)
        if os.path.exists(p):
            picture_fit(slide, p, x, y, w, h, border=True)
            return
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, PALE, BLUE2,
          width=1.4, dash="dash", radius=0.04)
    if h >= 1.3:
        note_h, label_y, label_h = 0.5, 0.8, h - 1.05
        note_y, note_size, label_size = 0.28, 10, 8.5
    else:
        note_h, label_y = 0.35, 0.55
        label_h = max(h - 0.75, 0.15)
        note_y, note_size, label_size = 0.18, 9, 7.5
    add_text(slide, note, x + 0.15, y + note_y, w - 0.3, note_h, note_size, True,
             BLUE, PP_ALIGN.CENTER)
    add_text(slide, "\n".join(filenames), x + 0.15, y + label_y, w - 0.3, label_h,
             label_size, False, MUTED, PP_ALIGN.CENTER, spacing=1.1)


def source(slide, text):
    add_text(slide, text, 0.55, 7.08, 12.2, 0.18, 6.5, False, MUTED)


def progress_bar(slide, current):
    x0, y, w, gap, h = 0.55, 0.99, 1.94, 0.13, 0.22
    for i, lab in enumerate(STAGES):
        x = x0 + i * (w + gap)
        if i < current:
            rrect(slide, x, y, w, h, BLUE, BLUE, radius=0.3)
            add_text(slide, lab, x, y + 0.02, w, 0.18, 8, True, WHITE, PP_ALIGN.CENTER)
        elif i == current:
            rrect(slide, x, y, w, h, BLUE2, BLUE, radius=0.3)
            add_text(slide, lab, x, y + 0.02, w, 0.18, 8.5, True, WHITE, PP_ALIGN.CENTER)
        else:
            rrect(slide, x, y, w, h, PALE, BORDER, radius=0.3)
            add_text(slide, lab, x, y + 0.02, w, 0.18, 8, False, MUTED, PP_ALIGN.CENTER)
        if i < len(STAGES) - 1:
            chevron(slide, x + w + 0.005, y + 0.03, 0.12, BORDER)


def header(slide, stage_idx, kicker, title, sub=""):
    rect(slide, 0, 0, 13.333, 0.06, BLUE)
    rect(slide, 0.55, 0.32, 1.5, 0.32, BLUE)
    add_text(slide, kicker, 0.55, 0.35, 1.5, 0.26, 10, True, WHITE, PP_ALIGN.CENTER)
    add_text(slide, title, 2.2, 0.28, 10.55, 0.4, 20, True, BLACK)
    if sub:
        add_text(slide, sub, 2.2, 0.72, 10.55, 0.26, 9.5, False, MUTED)
    progress_bar(slide, stage_idx)
    rect(slide, 0.55, 1.32, 12.23, 0.014, BORDER)


def conclusion(slide, text, color=BLUE):
    rrect(slide, 0.55, 1.44, 12.23, 0.32, PALE if color == BLUE else RED_PALE, BORDER, radius=0.14)
    add_text(slide, text, 0.75, 1.47, 11.85, 0.26, 10.5, True, color, PP_ALIGN.CENTER)


def page(slide, n):
    add_text(slide, "生成式文本隐写 | 科研面试汇报", 0.55, 7.08, 4.0, 0.16, 6.5, False, MUTED)
    add_text(slide, "%d / 10" % n, 12.35, 7.08, 0.8, 0.16, 7, False, MUTED, PP_ALIGN.RIGHT)


def parse_speech(path):
    notes = {}
    if not os.path.exists(path):
        return notes
    text = open(path, "r", encoding="utf-8").read()
    blocks = re.split(r"\n## Slide (\d+):", text)
    for i in range(1, len(blocks) - 1, 2):
        notes[int(blocks[i])] = blocks[i + 1].strip()
    return notes


def new_slide(prs, notes, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    page(s, n)
    if n in notes:
        s.notes_slide.notes_text_frame.text = notes[n]
    return s


def draw_bars(slide, x0, y0, items, bar_w=0.42, gap=0.42, max_h=1.7):
    for i, (label, v, color) in enumerate(items):
        x = x0 + i * (bar_w + gap)
        h = max_h * v
        rrect(slide, x, y0 + (max_h - h), bar_w, h, color, color, radius=0.08)
        add_text(slide, label, x - 0.08, y0 + max_h + 0.03, bar_w + 0.16, 0.16, 7.5,
                 False, MUTED, PP_ALIGN.CENTER)


def draw_bits(slide, x0, y0, bits, w=0.24, h=0.24, gap=0.07):
    for i, b in enumerate(bits):
        x = x0 + i * (w + gap)
        rrect(slide, x, y0, w, h, TEAL if b else WHITE, TEAL, radius=0.1)
        add_text(slide, str(b), x, y0 + 0.015, w, 0.18, 8, True, WHITE if b else TEAL,
                 PP_ALIGN.CENTER)


def big_number(slide, x, y, w, h, number, caption, accent=BLUE, fill=WHITE):
    rrect(slide, x, y, w, h, fill, BORDER)
    rect(slide, x, y, w, 0.045, accent)
    add_text(slide, number, x, y + 0.18, w, 0.62, 26, True, accent, PP_ALIGN.CENTER)
    add_text(slide, caption, x + 0.12, y + 0.9, w - 0.24, h - 1.05, 8.5, False, GRAY,
             PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, spacing=1.05)


def build():
    os.makedirs(ASSETS, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "大语言模型生成式文本隐写的可靠编码与跨精度回放研究"
    prs.core_properties.author = "项目答辩材料"
    prs.core_properties.subject = "科研面试 / 保研面试汇报（codex-ppt 工作流，插图待生图）"
    notes = parse_speech(SPEECH)

    # ---------- Slide 1 封面 ----------
    s = new_slide(prs, notes, 1)
    rect(s, 0, 0, 5.35, 7.5, BLUE)
    add_text(s, "科研面试汇报 · 保研面试", 0.6, 0.85, 4.3, 0.3, 11, True, RGBColor(0xBD, 0xD7, 0xF2))
    add_text(s, "大语言模型\n生成式文本隐写的\n可靠编码与\n跨精度回放研究", 0.6, 1.35, 4.35, 2.6, 23, True, WHITE, spacing=1.12)
    rect(s, 0.62, 4.15, 1.2, 0.03, RGBColor(0x9A, 0xC7, 0xF2))
    add_text(s, "复现 · 比较 · 扩展 SparSamp 与 RRC 两条路线", 0.6, 4.35, 4.3, 0.3, 10.5, False, RGBColor(0xD5, 0xE7, 0xF8))
    add_text(s, "姓名：____________\n学校 / 专业：____________\n时长：8 分钟", 0.6, 5.6, 4.3, 0.9, 10.5, False, WHITE, spacing=1.3)
    placeholder(s, 5.75, 1.45, 7.0, 4.35, ["fig_cover_stego_channel.png"])
    draw_bars(s, 6.35, 5.2, [("t1", 0.9, BORDER), ("t2", 1.5, TEAL), ("t3", 0.62, BORDER),
                              ("t4", 1.12, BORDER), ("t5", 0.42, BORDER), ("t6", 0.75, BORDER)])
    draw_bits(s, 6.35, 5.9, [0, 1, 0, 1, 1, 0, 1, 0])
    add_text(s, "状态：working paper / 项目阶段成果", 6.35, 6.55, 5.0, 0.25, 9.5, True, RED)
    source(s, "依据：README.md、RESEARCH_BRIEF.md。右侧为插图预留位与原生示意（非实验数据）。")

    # ---------- Slide 2 研究背景与动机 ----------
    s = new_slide(prs, notes, 2)
    header(s, 0, "01 背景", "为什么研究大语言模型文本隐写？")
    conclusion(s, "结论：可靠恢复与数值回放一致性，是生成式文本隐写落地的前提。")
    card(s, 0.55, 1.92, 2.85, 1.6, "普通生成", "每步输出 next-token 概率分布，据此采样。", BLUE, PALE)
    chevron(s, 3.5, 2.55, 0.45, BLUE2)
    card(s, 4.05, 1.92, 2.85, 1.6, "隐写编码", "共享密钥把秘密比特映射到概率空间。", TEAL, TEAL_PALE)
    chevron(s, 7.0, 2.55, 0.45, BLUE2)
    card(s, 7.55, 1.92, 2.6, 1.6, "接收端回放", "相同密钥/模型/配置反向恢复消息。", BLUE2, PALE)
    placeholder(s, 10.4, 1.92, 2.4, 1.6, ["fig_stego_channel.png", "（可复用封面图）"])
    card(s, 0.55, 3.75, 6.1, 1.25, "一个 token 的分叉会被放大",
         "FP16 与 BF16 的微小 logit 差异可能改变采样决策，并沿自回归轨迹持续传播。",
         RED, RED_PALE)
    card(s, 6.95, 3.75, 5.85, 1.25, "实证引子（R044）",
         "FP16→BF16 不做修正：60 条轨迹仅 10 条精确重放（16.7%）。", BLUE, PALE)
    add_text(s, "研究重点：自然性 · 容量 · 可恢复性 · 可检测性 · 数值稳定性（后两者为本项目重点）",
             0.55, 5.3, 12.23, 0.4, 10.5, True, BLACK, PP_ALIGN.CENTER)
    source(s, "依据：README.md；docs/reproducibility/R044_ANALYSIS.md。插图区为预留位（待生图）。")

    # ---------- Slide 3 研究问题与证据边界 ----------
    s = new_slide(prs, notes, 3)
    header(s, 1, "02 问题", "研究问题与证据边界")
    conclusion(s, "研究问题：公开整数概率合同 + 稀疏目标特定修正，能否实现精确回放？")
    items = [
        (0.55, 1.95, "自然性", "文本应像正常生成", BLUE),
        (6.85, 1.95, "容量", "每个 token 承载多少信息", ORANGE),
        (0.55, 4.15, "可恢复性", "接收端能否无误解码", TEAL),
        (6.85, 4.15, "可检测性", "观察者能否区分 stego 与 cover", RED),
    ]
    for x, y, hd, bd, ac in items:
        card(s, x, y, 5.9, 1.5, hd, bd, ac, WHITE, bsize=10)
    shape(s, MSO_SHAPE.OVAL, 5.45, 2.95, 2.45, 1.8, PALE, BLUE2)
    add_text(s, "项目聚焦\n可靠恢复\n+ 回放一致性", 5.6, 3.1, 2.15, 1.05, 12, True, BLUE,
             PP_ALIGN.CENTER, spacing=1.1)
    add_text(s, "关键边界：逐步匹配 token 分布 ≠ 变长全文分布匹配（精确反例 TV = 1/2）。",
             0.55, 5.85, 12.23, 0.32, 11, True, ORANGE, PP_ALIGN.CENTER)
    add_text(s, "不宣称：零 KL、完全不可检测、通用安全或跨硬件确定性。",
             0.55, 6.28, 12.23, 0.32, 11, True, RED, PP_ALIGN.CENTER)
    source(s, "依据：docs/verified_rrc_theory.md §3/§4、RESEARCH_BRIEF.md。")

    # ---------- Slide 4 技术路线总览 ----------
    s = new_slide(prs, notes, 4)
    header(s, 2, "03 路线", "两条生成式文本隐写路线：复现、比较、扩展")
    conclusion(s, "两条路线并列；官方 artifact 兼容复现通过核心门禁。")
    card(s, 0.55, 1.92, 5.95, 2.3, "路线 A：SparSamp（USENIX Security 2025）",
         "在候选 token 概率分布上做受控稀疏采样，block 收缩到单点后恢复 bit。\n项目角色：复现与比较对象。",
         BLUE, PALE, bsize=10)
    card(s, 6.85, 1.92, 5.95, 2.3, "路线 B：Range Coding / RRC（ACL 2026）",
         "把秘密消息映射到概率区间，逐 token 旋转并收缩；接收端反向重放。\n项目角色：可靠性审计与扩展对象。",
         TEAL, TEAL_PALE, bsize=10)
    chips = [
        ("R001", "105 tokens / 576 bits\n5.486 bit/token 精确解码"),
        ("R002", "1,193 / 1,200 完成"),
        ("R002", "846/846 无歧义\n精确解码"),
        ("R002", "16/16 容量误差 ≤5%\n最大 4.12%"),
    ]
    for i, (tag, txt) in enumerate(chips):
        x = 0.55 + i * 3.14
        rrect(s, x, 4.6, 2.85, 1.2, PALE, BORDER)
        add_text(s, tag, x + 0.16, 4.7, 1.0, 0.22, 9, True, BLUE2)
        add_text(s, txt, x + 0.16, 4.98, 2.55, 0.7, 9, False, GRAY, anchor=MSO_ANCHOR.TOP, spacing=1.1)
    add_text(s, "动作：复现、比较、扩展；不写“融合成新算法”。", 0.55, 6.1, 12.23, 0.3, 10.5,
             True, RED, PP_ALIGN.CENTER)
    source(s, "依据：docs/rrc.md、docs/reproducibility/R002_OFFICIAL_MATRIX.md、RESEARCH_BRIEF.md。")

    # ---------- Slide 5 SparSamp 原理与复现 ----------
    s = new_slide(prs, notes, 5)
    header(s, 2, "04 SparSamp", "SparSamp：概率分布上的“带密钥稀疏采样”")
    conclusion(s, "关键不是“强行指定词”，而是在模型原有概率分布上做可回放的受控采样。")
    steps = [
        ("1 概率分布", "模型给候选 token 分配概率，构成采样空间。"),
        ("2 密钥偏移", "共享密钥/PRF 生成受控偏移，双方可重放。"),
        ("3 消息落点", "消息块决定稀疏区间内的落点。"),
        ("4 收缩恢复", "block 收缩到单点后确定全部 bit。"),
    ]
    for i, (hd, bd) in enumerate(steps):
        y = 1.95 + i * 1.12
        rrect(s, 0.55, y, 5.9, 0.95, PALE if i % 2 == 0 else WHITE, BORDER)
        add_text(s, hd, 0.75, y + 0.08, 2.4, 0.24, 11.5, True, BLUE)
        add_text(s, bd, 0.75, y + 0.38, 5.4, 0.5, 9.5, False, GRAY, anchor=MSO_ANCHOR.TOP)
    placeholder(s, 6.85, 1.95, 5.95, 3.5, ["fig_sparsamp_sampling.png"])
    rrect(s, 6.85, 5.65, 5.95, 0.75, ORANGE_PALE, BORDER)
    add_text(s, "负面结果保留：FH-SparSamp v1 在 128-token 预算下成功率为 2/6，不作为推荐默认。",
             7.0, 5.72, 5.65, 0.6, 9, True, ORANGE, anchor=MSO_ANCHOR.TOP, spacing=1.12)
    source(s, "依据：docs/fh_sparsamp.md、refine-logs/R005_FH_V1_RESULTS.md。插图区为预留位（待生图）。")

    # ---------- Slide 6 RRC 原理与关键审计发现（逐步推导条） ----------
    s = new_slide(prs, notes, 6)
    header(s, 2, "05 RRC", "RRC 原理与关键审计发现：局部停止条件的跨边界问题")
    conclusion(s, "发现：局部条件 −1/2 < mid−d ≤ 1/2 不是精确恢复的充分条件（限定：本项目采用的 Algorithm 3/4 形式化）。",
               color=RED)
    steps = [("初始秘密点", "m ∈ [0, 2^l)"), ("旋转 + 概率切分", "o_t 旋转，按 token 概率分割"),
             ("选中子区间", "保留含秘密点的区间，逐 token 收缩")]
    for i, (hd, bd) in enumerate(steps):
        x = 0.55 + i * 4.15
        card(s, x, 1.9, 3.7, 1.1, hd, bd, BLUE2 if i != 1 else TEAL, PALE if i != 1 else TEAL_PALE,
             hsize=11, bsize=8.5)
        if i < 2:
            chevron(s, x + 3.8, 2.32, 0.3, BLUE2)
    add_text(s, "精确反例逐步推导（M=8，m=3）", 0.55, 3.2, 6.0, 0.28, 11.5, True, BLACK)
    nodes = [
        ("初始", "I⁻¹=[0,8)\nm=3", WHITE),
        ("t=0", "选中 I⁰=[3,5)", PALE),
        ("t=1", "o₁=1/2\nI¹=[3,9/2)", PALE),
        ("反向", "mid=15/4\n逆旋转得 19/4", ORANGE_PALE),
        ("结果", "round(19/4)=5\n5 ≠ 3", RED_PALE),
    ]
    nw, nh, gap = 2.1, 1.15, 0.24
    for i, (hd, bd, fl) in enumerate(nodes):
        x = 0.55 + i * (nw + gap)
        rrect(s, x, 3.55, nw, nh, fl, BORDER, radius=0.08)
        add_text(s, hd, x, 3.63, nw, 0.24, 9, True, RED if i == 4 else BLUE, PP_ALIGN.CENTER)
        add_text(s, bd, x, 3.92, nw, 0.7, 9, False, GRAY, PP_ALIGN.CENTER, spacing=1.05)
        if i < 4:
            chevron(s, x + nw + 0.01, 3.99, 0.22, BLUE2)
    add_text(s, "失败区域：o₀∈[0,1/64]，o₁∈[15/32,17/32] → 测度 1/1024（正概率，非孤立点）\n"
                 "根因：反向模旋转跨越切点——圆周距离保持 ≠ 线性距离保持",
             0.55, 5.0, 7.3, 0.85, 9.5, False, GRAY, anchor=MSO_ANCHOR.TOP, spacing=1.2)
    placeholder(s, 8.05, 4.95, 4.7, 0.95, ["fig_counterexample_wrap.png"])
    rrect(s, 0.55, 6.05, 12.23, 0.5, RED_PALE, BORDER)
    add_text(s, "必须真正执行接收端逆过程验证，而不是只检查局部距离。",
             0.75, 6.13, 11.85, 0.3, 10.5, True, RED, PP_ALIGN.CENTER)
    source(s, "依据：docs/rrc.md、docs/verified_rrc_theory.md §4、refine-logs/R019_VERIFIED_RRC_THEORY_AUDIT.md。")

    # ---------- Slide 7 方法改进一 ----------
    s = new_slide(prs, notes, 7)
    header(s, 3, "06 方法一", "Verified-RRC 与 Fixed-Length：确认能恢复才停止")
    conclusion(s, "固定公开长度 + HMAC 认证帧，缓解停止长度泄露。")
    flow = [("候选停止", "局部条件满足"), ("逆重放验证", "完整模拟接收端"), ("HMAC 认证", "验证消息前缀"),
            ("固定长度", "采样至公开 N")]
    for i, (hd, bd) in enumerate(flow):
        x = 0.55 + i * 3.14
        card(s, x, 1.95, 2.6, 1.05, hd, bd, BLUE2, PALE, hsize=10.5, bsize=8.5)
        if i < 3:
            chevron(s, x + 2.7, 2.34, 0.32, ORANGE)
    stats = [
        ("原局部规则", "Mock 精确恢复率 87%–96%", RED, RED_PALE),
        ("Verified-RRC", "500/500 精确恢复\n平均 +0.04–0.15 token", TEAL, TEAL_PALE),
        ("R020 固定长度", "Mock 400/400\n0 错 key 接受", BLUE, PALE),
    ]
    for i, (hd, bd, ac, fl) in enumerate(stats):
        x = 0.55 + i * 4.15
        card(s, x, 3.3, 3.75, 1.5, hd, bd, ac, fl, hsize=11, bsize=9)
    placeholder(s, 9.0, 3.3, 3.8, 1.5, ["fig_hmac_frame.png"])
    rrect(s, 0.55, 5.1, 12.23, 0.6, PALE, BORDER)
    add_text(s, "认证帧：F = m ‖ HMAC(K_auth, domain ‖ l ‖ context ‖ m)[:a]，默认 a=128",
             0.75, 5.2, 11.85, 0.3, 10.5, True, BLUE, PP_ALIGN.CENTER)
    add_text(s, "边界：非任意模型必然有限终止；R021 4 对 matched cover 未见稳定熵偏移（1.110 vs 1.124 bit/token），不构成不可检测结论。",
             0.55, 5.95, 12.23, 0.55, 9, True, ORANGE, anchor=MSO_ANCHOR.TOP, spacing=1.1)
    source(s, "依据：refine-logs/R012、R020、R021；docs/fixed_length_rrc.md。插图区为预留位（待生图）。")

    # ---------- Slide 8 方法改进二：BDS ----------
    s = new_slide(prs, notes, 8)
    header(s, 3, "07 方法二", "BDS：跨 FP16 / BF16 只记录真正不稳定的决策")
    conclusion(s, "公开整数概率合同 + 有限枚举可行决策集；稳定步省略，不稳定步记录。")
    if os.path.exists(ARCH):
        picture_fit(s, ARCH, 0.55, 1.95, 5.9, 3.7, border=True)
        add_text(s, "工作论文 Figure 1（BDS 架构，既有图）", 0.55, 5.75, 5.9, 0.25, 8, True, MUTED,
                 PP_ALIGN.CENTER)
    else:
        placeholder(s, 0.55, 1.95, 5.9, 3.7, ["fig_bds_replay.png"])
    card(s, 6.85, 1.95, 5.95, 1.4, "问题（R044）",
         "FP16→BF16 未校正仅 10/60 精确；平均修正率 2.16% [1.80, 2.53]。", RED, RED_PALE)
    card(s, 6.85, 3.5, 5.95, 1.4, "机制（BDS）",
         "量化 bin / top-16 envelope / top-2 整数合同；在 bin 扰动半径内枚举可行决策。", TEAL, TEAL_PALE)
    card(s, 6.85, 5.05, 5.95, 1.4, "实证（R059）",
         "独立 seed 40/40 精确回放、0 false-safe；证书负载/完整轨迹 ≈54.6%。", BLUE, PALE)
    add_text(s, "新概念图预留位：fig_bds_replay.png（待生图，届时替换左图）",
             6.85, 6.5, 5.95, 0.25, 8, True, MUTED, PP_ALIGN.CENTER)
    source(s, "依据：docs/reproducibility/R055、R056、R058、R059；RESEARCH_BRIEF.md。边界：Qwen2.5-1.5B-Instruct、RTX 3060 Laptop、指定软件栈；不等于跨硬件绝对可复现。")

    # ---------- Slide 9 实验结果汇总（大数字卡） ----------
    s = new_slide(prs, notes, 9)
    header(s, 4, "08 结果", "实验结果汇总（真实数据）")
    conclusion(s, "主实验：校正后 60/60 精确重放，平均修正率 2.16% [1.80, 2.53]。")
    if os.path.exists(FIG2):
        picture_fit(s, FIG2, 0.55, 1.95, 6.3, 4.15, border=True)
        add_text(s, "R044：校正后 60/60 vs 未校正 10/60（真实数据）",
                 0.55, 6.18, 6.3, 0.25, 8.5, True, MUTED, PP_ALIGN.CENTER)
    else:
        placeholder(s, 0.55, 1.95, 6.3, 4.15, ["fig2_main_scale.png"])
    nums = [
        ("60/60", "校正后精确重放（R044）", BLUE, WHITE),
        ("10/60", "未校正对照组", RED, WHITE),
        ("2.16%", "平均修正率 [1.80, 2.53]", TEAL, WHITE),
        ("24.76%", "SPRC 包/完整轨迹\n（1,148 B vs 4,636 B）", BLUE2, WHITE),
    ]
    for i, (num, cap, ac, fl) in enumerate(nums):
        x = 7.1 + (i % 2) * 3.1
        y = 1.95 + (i // 2) * 2.15
        big_number(s, x, y, 2.9, 1.95, num, cap, ac, fl)
    add_text(s, "补充：反向精度方向 BF16→FP16 20/20（R046）；口径：payload-only 6.65% / referenced 24.76% / self-contained 63.89% 是不同估计量。",
             0.55, 6.5, 12.23, 0.3, 9, True, ORANGE, PP_ALIGN.CENTER)
    source(s, "依据：figures/source_data/figure_02_source.csv；docs/reproducibility/R044_ANALYSIS.md、R046_ABLATION_ANALYSIS.md、R051_REPLAY_BASELINES.md。")

    # ---------- Slide 10 局限、未来与总结 ----------
    s = new_slide(prs, notes, 10)
    header(s, 5, "09 总结", "局限、未来与总结")
    conclusion(s, "结论：先保证“能正确回来”，再讨论“能否稳定重放”；所有结论均为阶段性、条件化的证据。")
    res = [
        ("总结 1", "RRC 停止正确性审计 + Verified-RRC（Mock 500/500）", TEAL, TEAL_PALE),
        ("总结 2", "跨精度稀疏修正回放：校正后 60/60 vs 未校正 10/60", BLUE, PALE),
        ("总结 3", "BDS 独立确认：40/40 精确回放、0 false-safe（R059）", BLUE, PALE),
    ]
    for i, (hd, bd, ac, fl) in enumerate(res):
        x = 0.55 + i * 4.15
        card(s, x, 1.95, 3.75, 1.6, hd, bd, ac, fl, hsize=11, bsize=9)
    card(s, 0.55, 3.85, 6.05, 1.5, "局限（如实声明）",
         "单模型单 GPU；无真人盲评（R048 未收集）；无跨硬件验证；无任意模型必然终止证明。",
         ORANGE, ORANGE_PALE)
    card(s, 6.85, 3.85, 6.0, 1.5, "未来工作",
         "扩大模型/硬件矩阵；固定长度失败率区间；整数频数合同；盲评与检测实验。", BLUE, PALE)
    add_text(s, "状态：working paper / 项目阶段成果；不宣称已发表、通用安全、完全不可检测、跨硬件保证或零 KL。",
             0.55, 5.75, 12.23, 0.35, 11, True, RED, PP_ALIGN.CENTER)
    source(s, "依据：RESEARCH_BRIEF.md、paper/CLAIM_EVIDENCE_MAP.md、paper/AUTHOR_NOTES_ZH.md。")

    prs.save(OUT)
    print("saved:", OUT)
    print("slides:", len(prs.slides._sldIdLst))


if __name__ == "__main__":
    build()
