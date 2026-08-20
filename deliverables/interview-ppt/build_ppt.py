# -*- coding: utf-8 -*-
"""Build the 10-slide 16:9 interview PPT (native editable slides + speaker notes).

Usage:
  python build_ppt.py

The script reads speech.md for per-slide notes and embeds generated illustrations
from assets/ when present; otherwise it draws native editable diagrams and marks
them as 示意. Real data figures are embedded from the repo's deliverables/assets.
"""

import os
import re

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(ROOT, "assets")
OUT = os.path.join(ROOT, "生成式文本隐写_科研面试汇报_v2.pptx")
SPEECH = os.path.join(ROOT, "speech.md")

REPO = r"C:\Users\41462\Documents\隐写"
REAL_FIG2 = os.path.join(REPO, "deliverables", "assets", "fig2_main_scale.png")
REAL_FIG4 = os.path.join(REPO, "deliverables", "assets", "fig4_precision_direction.png")
REAL_ARCH = os.path.join(REPO, "paper", "ccfa_bounded_decision", "figures", "figure_01_architecture.png")

NAVY = RGBColor(0x17, 0x3B, 0x64)
BLUE = RGBColor(0x24, 0x5C, 0x97)
CYAN = RGBColor(0x0C, 0x8F, 0x88)
ORANGE = RGBColor(0xD8, 0x78, 0x1F)
RED = RGBColor(0xC2, 0x48, 0x3B)
TEXT = RGBColor(0x27, 0x31, 0x3D)
MUTED = RGBColor(0x66, 0x70, 0x85)
PALE = RGBColor(0xED, 0xF3, 0xF8)
GREEN_PALE = RGBColor(0xE8, 0xF5, 0xF1)
ORANGE_PALE = RGBColor(0xFF, 0xF1, 0xE5)
RED_PALE = RGBColor(0xFD, 0xEE, 0xEE)
LINE = RGBColor(0xD7, 0xDE, 0xE7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xFA, 0xFB, 0xFC)
BAR = RGBColor(0xAF, 0xC6, 0xD9)

FONT = "Microsoft YaHei"


def set_run(run, size, bold=False, color=TEXT, italic=False, font=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", font)


def add_text(slide, text, x, y, w, h, size=13, bold=False, color=TEXT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, italic=False,
             wrap=True, spacing=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing is not None:
            p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        set_run(r, size, bold, color, italic)
    return box


def add_shape(slide, shape_type, x, y, w, h, fill=WHITE, line=LINE, radius=None):
    sp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(0.8)
    sp.shadow.inherit = False
    if radius is not None and shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def rect(slide, x, y, w, h, fill, line=None):
    return add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line)


def rrect(slide, x, y, w, h, fill, line, radius=0.07):
    return add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, radius)


def chevron(slide, x, y, w, color=BLUE):
    return add_shape(slide, MSO_SHAPE.CHEVRON, x, y, w, 0.24, color, color)


def card(slide, x, y, w, h, heading, body, accent=BLUE, fill=WHITE, hsize=13, bsize=10.5):
    rrect(slide, x, y, w, h, fill, LINE)
    rect(slide, x, y, 0.06, h, accent)
    add_text(slide, heading, x + 0.22, y + 0.14, w - 0.42, 0.32, hsize, True, accent)
    add_text(slide, body, x + 0.22, y + 0.52, w - 0.42, h - 0.68, bsize, False, TEXT,
             anchor=MSO_ANCHOR.TOP, spacing=1.12)


def pill(slide, text, x, y, w, color, pale, size=8):
    rrect(slide, x, y, w, 0.31, pale, pale, radius=0.5)
    add_text(slide, text, x + 0.05, y + 0.055, w - 0.1, 0.2, size, True, color,
             align=PP_ALIGN.CENTER)


def source(slide, text):
    add_text(slide, text, 0.62, 7.08, 11.9, 0.2, 6.7, False, MUTED)


def header(slide, kicker, title, sub=""):
    rect(slide, 0, 0, 13.333, 0.11, NAVY)
    add_text(slide, kicker, 0.62, 0.30, 6.0, 0.24, 9, True, BLUE)
    add_text(slide, title, 0.62, 0.60, 12.1, 0.55, 23, True, NAVY)
    if sub:
        add_text(slide, sub, 0.64, 1.22, 12.1, 0.30, 10, False, MUTED)


def footer(slide, n):
    add_text(slide, "生成式文本隐写 | 科研面试汇报", 0.58, 7.08, 4.2, 0.18, 6.8, False, MUTED)
    add_text(slide, "%d / 10" % n, 12.25, 7.08, 0.9, 0.18, 7, False, MUTED, PP_ALIGN.RIGHT)


def img_path(name):
    p = os.path.join(ASSETS, name)
    return p if os.path.exists(p) else None


def add_picture_fit(slide, path, x, y, max_w, max_h):
    from PIL import Image
    with Image.open(path) as im:
        w, h = im.size
    ratio = min(max_w / w, max_h / h)
    dw, dh = w * ratio, h * ratio
    slide.shapes.add_picture(path, Inches(x + (max_w - dw) / 2), Inches(y + (max_h - dh) / 2),
                             Inches(dw), Inches(dh))


def frame_badge(slide, x, y, text="示意（非实验数据）"):
    pill(slide, text, x, y, 2.2, ORANGE, ORANGE_PALE, size=7.5)


def parse_speech(path):
    notes = {}
    if not os.path.exists(path):
        return notes
    text = open(path, "r", encoding="utf-8").read()
    blocks = re.split(r"\n## Slide (\d+):", text)
    # blocks[0] is preamble; then pairs (num, body)
    for i in range(1, len(blocks) - 1, 2):
        num = int(blocks[i])
        body = blocks[i + 1].strip()
        notes[num] = body
    return notes


def new_slide(prs, notes, n):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    footer(slide, n)
    if n in notes:
        slide.notes_slide.notes_text_frame.text = notes[n]
    return slide


def draw_probability_bars(slide, x0, y0, items, bar_w=0.36, gap=0.36, max_h=1.9):
    for i, (label, v, color) in enumerate(items):
        x = x0 + i * (bar_w + gap)
        h = max_h * v
        rrect(slide, x, y0 + (max_h - h), bar_w, h, color, color, radius=0.05)
        add_text(slide, label, x - 0.06, y0 + max_h + 0.04, bar_w + 0.12, 0.18, 8.5,
                 False, MUTED, PP_ALIGN.CENTER)


def draw_bits(slide, x0, y0, bits, w=0.30, h=0.30, gap=0.08):
    for i, b in enumerate(bits):
        x = x0 + i * (w + gap)
        rrect(slide, x, y0, w, h, CYAN if b else WHITE, CYAN, radius=0.08)
        add_text(slide, str(b), x, y0 + 0.02, w, 0.2, 9, True, WHITE if b else CYAN,
                 PP_ALIGN.CENTER)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "大语言模型生成式文本隐写的可靠编码与跨精度回放研究"
    prs.core_properties.author = "项目答辩材料"
    prs.core_properties.subject = "科研面试 / 保研面试汇报"
    notes = parse_speech(SPEECH)

    # ---------- Slide 1: Cover ----------
    s = new_slide(prs, notes, 1)
    rect(s, 0, 0, 13.333, 7.5, BG)
    rect(s, 0, 0, 5.45, 7.5, NAVY)
    add_text(s, "科研面试汇报 · 保研面试", 0.62, 0.85, 4.2, 0.30, 11, True, RGBColor(0x9A, 0xD7, 0xD0))
    add_text(s, "大语言模型\n生成式文本隐写的\n可靠编码与\n跨精度回放研究", 0.62, 1.35, 4.4, 2.6, 24, True, WHITE, spacing=1.12)
    add_text(s, "复现、比较并扩展 SparSamp 与 RRC 两条路线", 0.66, 4.05, 4.3, 0.30, 11, False, RGBColor(0xD5, 0xE7, 0xF5))
    add_text(s, "姓名：____________\n学校 / 专业：____________\n汇报时长：8 分钟", 0.66, 5.55, 4.3, 1.0, 10.5, False, WHITE, spacing=1.25)
    fig = img_path("fig_cover_stego_channel.png")
    if fig:
        add_picture_fit(s, fig, 5.75, 1.35, 7.0, 4.6)
    else:
        add_text(s, "概率空间中的秘密信息编码（示意）", 5.95, 1.25, 4.6, 0.3, 12, True, NAVY)
        items = [("t1", 0.9, BAR), ("t2", 1.5, CYAN), ("t3", 0.62, BAR), ("t4", 1.12, BAR),
                 ("t5", 0.42, BAR), ("t6", 0.75, BAR), ("t7", 0.3, BAR)]
        draw_probability_bars(s, 6.15, 4.35, items)
        draw_bits(s, 6.15, 5.35, [0, 1, 0, 1, 1, 0, 1, 0])
        chevron(s, 8.05, 5.15, 0.7, ORANGE)
        pill(s, "模型分布", 6.0, 6.05, 1.35, BLUE, PALE)
        pill(s, "密钥控制", 7.55, 6.05, 1.35, CYAN, GREEN_PALE)
        pill(s, "秘密比特", 9.1, 6.05, 1.35, ORANGE, ORANGE_PALE)
        frame_badge(s, 10.85, 6.05, text="示意（非实验数据）")
    source(s, "说明：右侧为本页概念示意，不代表实验数据；姓名/学校/专业为占位符，请替换。")

    # ---------- Slide 2: Background ----------
    s = new_slide(prs, notes, 2)
    header(s, "01 研究背景", "为什么研究大语言模型文本隐写？",
           "目标：在正常生成的文本中传递秘密信息，同时尽量保持生成分布与文本质量。")
    card(s, 0.68, 1.85, 2.6, 3.6, "普通生成", "语言模型给每个候选 token 分配概率，并据此采样下一个词。", BLUE, PALE)
    tokens = [("天气", 0.88), ("今天", 0.65), ("适合", 0.42), ("学习", 0.28)]
    for i, (t, v) in enumerate(tokens):
        y = 2.75 + i * 0.58
        add_text(s, t, 1.0, y, 0.55, 0.2, 10)
        rrect(s, 1.66, y + 0.02, 1.5 * v, 0.15, BLUE if i == 0 else BAR,
              BLUE if i == 0 else BAR, radius=0.2)
    chevron(s, 3.45, 3.45, 0.8, ORANGE)
    card(s, 4.45, 1.85, 3.5, 3.6, "隐写编码", "共享密钥把秘密比特映射到概率空间，让选出的 token 同时承担“自然语言”与“信息载体”。", CYAN, GREEN_PALE)
    draw_bits(s, 5.0, 3.05, [1, 0, 1, 1, 0, 0], w=0.22, h=0.22, gap=0.08)
    chevron(s, 6.05, 3.45, 0.88, ORANGE)
    card(s, 8.18, 1.85, 4.36, 3.6, "接收端回放", "接收端以相同模型、上下文、密钥和离散决策规则回放，恢复嵌入的消息。", NAVY, PALE)
    add_text(s, "研究重点", 0.72, 5.75, 1.1, 0.2, 10, True, MUTED)
    pill(s, "自然性", 1.95, 5.72, 1.28, BLUE, PALE)
    pill(s, "可恢复性", 3.45, 5.72, 1.28, CYAN, GREEN_PALE)
    pill(s, "长度侧信道", 4.95, 5.72, 1.28, ORANGE, ORANGE_PALE)
    pill(s, "数值稳定性", 6.45, 5.72, 1.28, RED, RED_PALE)
    add_text(s, "实证背景（R044）：FP16 → BF16 不做修正时，60 条轨迹只有 10 条能精确重放。",
             8.18, 5.72, 4.4, 0.5, 9.5, True, RED, anchor=MSO_ANCHOR.TOP)
    source(s, "概念依据：项目 README、docs/algorithm.md、docs/reproducibility/R044_ANALYSIS.md；概率条为示意。")

    # ---------- Slide 3: Problem definition ----------
    s = new_slide(prs, notes, 3)
    header(s, "02 问题定义", "四个目标互相牵制，不能只看其中一个",
           "本项目优先审计“可靠恢复”与“数值回放一致性”，其他目标只报告证据边界。")
    items = [
        (1.0, 1.85, "自然性", "文本应像正常生成", BLUE, PALE),
        (7.05, 1.85, "容量", "每个 token 能承载多少信息", ORANGE, ORANGE_PALE),
        (1.0, 4.15, "可恢复性", "接收端能否无误解码", CYAN, GREEN_PALE),
        (7.05, 4.15, "可检测性", "观察者能否区分 stego 与 cover", RED, RED_PALE),
    ]
    for x, y, hd, bd, ac, fl in items:
        card(s, x, y, 5.2, 1.55, hd, bd, ac, fl)
    add_shape(s, MSO_SHAPE.OVAL, 5.42, 2.9, 2.5, 1.4, PALE, NAVY)
    add_text(s, "项目聚焦\n可靠恢复\n+ 回放一致性", 5.56, 3.05, 2.22, 0.85, 12, True, NAVY, PP_ALIGN.CENTER, spacing=1.1)
    add_text(s, "关键边界：逐步匹配 token 分布 ≠ 变长全文分布匹配（精确反例 TV = 1/2）。",
             1.06, 5.95, 11.2, 0.3, 10.5, True, ORANGE, PP_ALIGN.CENTER)
    add_text(s, "不宣称：通用安全证明、完全不可检测、零 KL 或跨硬件保证。",
             1.06, 6.35, 11.2, 0.3, 10.5, True, RED, PP_ALIGN.CENTER)
    source(s, "证据边界：docs/verified_rrc_theory.md、RESEARCH_BRIEF.md。")

    # ---------- Slide 4: Two routes ----------
    s = new_slide(prs, notes, 4)
    header(s, "03 技术路线", "复现、比较并扩展 SparSamp 与 RRC 两条路线",
           "两条路线并列；Verified-RRC、Fixed-Length RRC、BDS 是对 RRC 路线的可靠性扩展。")
    card(s, 0.72, 1.72, 5.5, 3.9, "路线 A：SparSamp（USENIX Security 2025）",
         "官方 artifact 兼容复现（Zenodo 15025436）：\n"
         "· Basic Test：105 tokens / 576 bits / 5.486 bit/token，精确解码\n"
         "· R002 矩阵：1,193/1,200 完成；无歧义子集 846/846 精确解码\n"
         "· 16/16 容量对比误差 ≤5%（最大 4.12%）",
         BLUE, PALE, hsize=12.5, bsize=9.5)
    add_text(s, "边界：兼容复现，非严格 Torch 2.2.2 环境复现；速度仅作硬件相关描述。",
             0.95, 5.28, 5.1, 0.5, 8.5, True, ORANGE, anchor=MSO_ANCHOR.TOP)
    card(s, 7.1, 1.72, 5.5, 3.9, "路线 B：Range Coding / RRC（ACL 2026）",
         "独立 clean-room 复现论文 Algorithm 3/4：\n"
         "· 秘密消息映射到概率区间，逐 token 旋转并收缩\n"
         "· 本项目审计对象：终止正确性、固定长度、跨精度回放",
         CYAN, GREEN_PALE, hsize=12.5, bsize=9.5)
    add_text(s, "论文局部停止条件的审计结论见第 7 页；表述限定在“本项目采用的形式化”内。",
             7.32, 5.28, 5.1, 0.5, 8.5, True, ORANGE, anchor=MSO_ANCHOR.TOP)
    chevron(s, 6.38, 3.45, 0.5, ORANGE)
    pill(s, "复现", 2.0, 5.95, 1.08, BLUE, PALE)
    pill(s, "比较", 3.95, 5.95, 1.08, ORANGE, ORANGE_PALE)
    pill(s, "可靠性扩展", 8.5, 5.95, 1.6, CYAN, GREEN_PALE)
    add_text(s, "不宣称：将 SparSamp 与 RRC 融合为一个新算法。", 2.55, 6.42, 8.2, 0.24, 10.5, True, RED, PP_ALIGN.CENTER)
    source(s, "参考：SparSamp（arXiv:2503.19499）；Yan & Murawaki, ACL 2026；docs/rrc.md、docs/reproducibility/R002_OFFICIAL_MATRIX.md。")

    # ---------- Slide 5: SparSamp ----------
    s = new_slide(prs, notes, 5)
    header(s, "04 SparSamp", "用概率分布做“带密钥的稀疏采样”",
           "下图为简化示意；真实实现还需共享配置、离散化规则与严格回放条件。")
    fig = img_path("fig_sparsamp_sampling.png")
    if fig:
        add_picture_fit(s, fig, 6.9, 1.55, 5.9, 4.3)
        frame_badge(s, 11.0, 5.95)
    else:
        add_text(s, "候选 token 概率", 0.85, 1.7, 2.2, 0.25, 12, True, NAVY)
        items = [("安全", 0.82, BLUE), ("可靠", 0.61, CYAN), ("高效", 0.40, BAR), ("自然", 0.25, BAR)]
        for i, (t, v, c) in enumerate(items):
            y = 2.25 + i * 0.62
            add_text(s, t, 1.0, y, 0.6, 0.2, 10.5)
            rrect(s, 1.75, y + 0.01, 2.2 * v, 0.18, c, c, radius=0.2)
            add_text(s, "%d%%" % int(v * 100), 4.1, y, 0.45, 0.18, 9, False, MUTED, PP_ALIGN.RIGHT)
        chevron(s, 4.62, 3.2, 0.8, ORANGE)
        card(s, 5.6, 1.85, 2.3, 2.6, "密钥 + 消息块", "共享密钥生成受控偏移；消息块决定候选空间内的落点。", ORANGE, ORANGE_PALE, bsize=9)
        chevron(s, 8.05, 3.2, 0.8, ORANGE)
        card(s, 9.0, 1.85, 3.55, 2.6, "稀疏采样选择 token", "在保持概率结构的前提下选出 token，并使接收端能重放同一决策。", CYAN, GREEN_PALE, bsize=9)
        pill(s, "概率分布", 1.2, 5.35, 1.35, BLUE, PALE)
        pill(s, "受控偏移", 5.9, 5.35, 1.35, ORANGE, ORANGE_PALE)
        pill(s, "稀疏决策", 9.7, 5.35, 1.35, CYAN, GREEN_PALE)
        frame_badge(s, 10.9, 6.1)
    add_text(s, "面试表达：SparSamp 的关键不是“强行指定词”，而是在模型原有概率分布上做可回放的受控采样。",
             0.9, 6.05, 11.6, 0.4, 10.5, False, TEXT, PP_ALIGN.CENTER)
    add_text(s, "本项目角色：复现与对比（FH-SparSamp v1 已被负面结果拒绝，不作为推荐默认）。",
             0.9, 6.45, 11.6, 0.3, 9, True, ORANGE, PP_ALIGN.CENTER)
    source(s, "概念依据：docs/fh_sparsamp.md、RESEARCH_BRIEF.md；本页示意图非实验数据。")

    # ---------- Slide 6: RRC ----------
    s = new_slide(prs, notes, 6)
    header(s, "05 RRC", "把秘密消息映射为区间，再逐 token 旋转收缩",
           "Range Coding 在项目中简称 RRC：秘密点 + 共享偏移旋转 + 概率分割 + 区间收缩。")
    heads = ["初始秘密点", "旋转并按概率切分", "选中子区间"]
    bodies = ["消息 m 对应初始区间 [0, 2^l) 内的一点。",
              "共享偏移 o_t 旋转秘密点；模型概率决定分割比例。",
              "输出 token 后保留含秘密点的子区间，继续下一步。"]
    for i in range(3):
        x = 1.0 + i * 4.1
        card(s, x, 1.7, 3.15, 2.55, heads[i], bodies[i],
             ORANGE if i == 1 else (CYAN if i == 2 else BLUE),
             ORANGE_PALE if i == 1 else (GREEN_PALE if i == 2 else PALE), bsize=9.5)
        if i < 2:
            chevron(s, x + 3.28, 2.85, 0.42, ORANGE)
    fig = img_path("fig_rrc_interval.png")
    if fig:
        add_picture_fit(s, fig, 0.95, 4.45, 11.5, 2.2)
        frame_badge(s, 10.85, 4.75)
    else:
        for i, (x0, frac, color, lab) in enumerate([
                (1.42, 0.45, BLUE, "I⁻¹ = [0, 2ˡ)"),
                (5.05, 0.33, ORANGE, "dᵗ = rotate(dᵗ⁻¹)"),
                (8.75, 0.2, CYAN, "Iᵗ ⊂ Iᵗ⁻¹")]):
            rect(s, x0, 5.0, 2.35, 0.16, BAR, None)
            rect(s, x0 + 0.3, 5.0, 2.35 * frac, 0.16, color, None)
            add_text(s, lab, x0, 5.35, 2.5, 0.2, 9.5, True, NAVY, PP_ALIGN.CENTER)
    add_text(s, "接收端按相同 token 前缀、偏移和分割规则反向重放，尝试恢复 m。",
             1.2, 6.6, 11.0, 0.3, 10.5, False, TEXT, PP_ALIGN.CENTER)
    source(s, "依据：docs/rrc.md、docs/verified_rrc_theory.md；示意不按具体模型概率绘制。")

    # ---------- Slide 7: Key finding ----------
    s = new_slide(prs, notes, 7)
    header(s, "06 关键发现", "局部“接近中点”并不总能保证正确恢复",
           "在本项目采用的 Algorithm 3/4 形式化下，跨越旋转切点时，局部停止条件不是精确恢复的充分条件。")
    steps = [
        ("t=0", "I⁻¹=[0,8)，m=3", "选中 I⁰=[3,5)", BLUE),
        ("t=1", "偏移 o₁=1/2", "选中 I¹=[3,9/2)", ORANGE),
        ("反向", "mid(I¹)−d¹=−1/4", "round(19/4)=5 ≠ 3", RED),
    ]
    for i, (hd, b1, b2, ac) in enumerate(steps):
        x = 0.85 + i * 4.15
        card(s, x, 1.72, 3.6, 1.95, hd, b1 + "\n" + b2, ac,
             RED_PALE if ac == RED else (ORANGE_PALE if ac == ORANGE else PALE), bsize=10)
        if i < 2:
            chevron(s, x + 3.7, 2.5, 0.38, ORANGE)
    fig = img_path("fig_counterexample_wrap.png")
    if fig:
        add_picture_fit(s, fig, 0.95, 3.95, 11.5, 2.15)
        frame_badge(s, 10.85, 4.35)
    else:
        add_text(s, "区间演化（示意）", 0.95, 3.95, 2.2, 0.22, 11, True, NAVY)
        rect(s, 1.35, 4.75, 9.9, 0.03, NAVY, None)
        marks = [(0, 1.35), (3, 4.93), (4.5, 6.35), (5, 6.62), (8, 9.47)]
        for val, x in marks:
            rect(s, x, 4.66, 0.03, 0.2, NAVY, None)
            add_text(s, str(val), x - 0.08, 4.92, 0.2, 0.16, 8, False, MUTED, PP_ALIGN.CENTER)
        rect(s, 4.93, 4.52, 2.72, 0.14, BLUE, None)
        rect(s, 4.93, 4.98, 1.69, 0.14, ORANGE, None)
        add_text(s, "I⁰", 5.9, 4.28, 0.5, 0.18, 9, True, BLUE, PP_ALIGN.CENTER)
        add_text(s, "I¹", 5.5, 5.22, 0.5, 0.18, 9, True, ORANGE, PP_ALIGN.CENTER)
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 2.1, 6.28, 9.12, 0.4, RED_PALE, RGBColor(0xF3, 0xC4, 0xBE), 0.15)
    add_text(s, "结论：不能只检查局部距离；必须真正执行接收端的逆过程验证。",
             2.25, 6.38, 8.82, 0.2, 10, True, RED, PP_ALIGN.CENTER)
    source(s, "精确有理数反例：docs/verified_rrc_theory.md 第 4 节；可用 scripts/verify_rrc_theory.py 重现。失败区域测度 1/1024。")

    # ---------- Slide 8: Improvements ----------
    s = new_slide(prs, notes, 8)
    header(s, "07 方法改进", "Verified-RRC：确认能恢复才停止；Fixed-Length + HMAC 缓解长度泄露",
           "R012 与 R020/R021 的测试结果如下；均带 Mock 或小样本范围限定。")
    flow = [("候选停止", "局部条件满足", ORANGE, ORANGE_PALE),
            ("逆重放验证", "完整模拟接收端", CYAN, GREEN_PALE),
            ("HMAC 认证", "验证消息前缀", BLUE, PALE),
            ("固定长度填充", "采样至公开 N", NAVY, PALE)]
    for i, (hd, bd, ac, fl) in enumerate(flow):
        x = 0.68 + i * 3.08
        card(s, x, 1.6, 2.6, 1.35, hd, bd, ac, fl, hsize=11.5, bsize=9)
        if i < 3:
            chevron(s, x + 2.7, 2.15, 0.28, ORANGE)
    card(s, 0.83, 3.35, 3.35, 1.75, "原局部停止规则",
         "Mock 精确恢复率\n87%–96%", RED, RED_PALE)
    card(s, 4.99, 3.35, 3.35, 1.75, "Verified-RRC",
         "500/500 精确恢复\n平均额外 +0.04–0.15 token", CYAN, GREEN_PALE)
    card(s, 9.15, 3.35, 3.35, 1.75, "代价边界",
         "单样本最多 +2 token\n≠ 任意模型必然终止", ORANGE, ORANGE_PALE)
    add_text(s, "R020：Mock 400/400 恢复、0 错 key 接受；Qwen smoke N=224 完成 4/4（小样本）。",
             0.83, 5.25, 8.0, 0.4, 9.5, True, NAVY)
    add_text(s, "认证帧：F = m ‖ HMAC(K_auth, domain ‖ l ‖ context ‖ m)[:a]，默认 a=128",
             0.95, 5.75, 11.4, 0.3, 10.5, True, NAVY)
    fig = img_path("fig_hmac_frame.png")
    if fig:
        add_picture_fit(s, fig, 9.0, 5.15, 3.9, 1.4)
        frame_badge(s, 10.0, 6.5)
    add_text(s, "边界：非任意模型必然有限终止；4 对 matched-cover 样本未见稳定熵偏移（1.110 vs 1.124 bit/token），不构成不可检测结论。",
             0.83, 6.25, 11.7, 0.55, 9, True, ORANGE, anchor=MSO_ANCHOR.TOP)
    source(s, "真实结果：refine-logs/R012_VERIFIED_RRC_RESULTS.md、R020_FIXED_LENGTH_RRC_V1.md、R021_FIXED_LENGTH_MATCHED_COVER.md。")

    # ---------- Slide 9: BDS ----------
    s = new_slide(prs, notes, 9)
    header(s, "08 工程可靠性", "BDS：跨 FP16 / BF16 时，只记录真正不稳定的决策",
           "微小 logit 差异可能改变一个 token，并在自回归生成中持续放大。")
    fig = img_path("fig_bds_replay.png")
    if fig:
        add_picture_fit(s, fig, 0.7, 1.6, 5.7, 4.3)
        frame_badge(s, 4.9, 6.0)
    else:
        add_picture_fit(s, REAL_ARCH, 0.7, 1.6, 5.7, 4.3)
        frame_badge(s, 4.9, 6.0)
    card(s, 6.75, 1.6, 5.85, 1.55, "问题（R044）",
         "FP16 → BF16 未校正仅 10/60 精确；平均修正率 2.16% [1.80, 2.53]。",
         RED, RED_PALE, bsize=9.5)
    card(s, 6.75, 3.3, 5.85, 1.55, "机制（BDS）",
         "公开整数概率合同 + 有限枚举可行决策集；稳定步省略，不稳定步记录稀疏修正。",
         CYAN, GREEN_PALE, bsize=9.5)
    card(s, 6.75, 5.0, 5.85, 1.55, "独立确认（R059）",
         "独立 seed 40/40 精确回放、0 false-safe；证书负载/完整轨迹 ≈54.6%。",
         BLUE, PALE, bsize=9.5)
    add_text(s, "边界：Qwen2.5-1.5B-Instruct、RTX 3060 Laptop、指定软件栈；不等于跨硬件绝对可复现。",
             0.75, 6.5, 11.9, 0.3, 10, True, ORANGE, PP_ALIGN.CENTER)
    source(s, "真实结果：docs/reproducibility/R058_BOUNDED_DECISION_SET_RESULTS.md、R059_INDEPENDENT_BOUNDED_DECISION_RESULTS.md；左侧为概念图/工作论文 Figure 1。")

    # ---------- Slide 10: Summary ----------
    s = new_slide(prs, notes, 10)
    header(s, "09 结果与总结", "先保证“能正确回来”，再讨论“能否稳定重放”",
           "阶段性 working paper / 项目阶段成果；不代表已发表、已投稿或通用安全结论。")
    if os.path.exists(REAL_FIG2):
        add_picture_fit(s, REAL_FIG2, 0.62, 1.62, 6.2, 4.6)
        add_text(s, "真实数据（R044）：校正后 60/60，未校正 10/60",
                 0.62, 6.35, 6.2, 0.25, 9, True, MUTED, PP_ALIGN.CENTER)
    card(s, 7.15, 1.62, 5.45, 1.2, "结果 1：恢复可靠性",
         "RRC 停止审计 + Verified-RRC：Mock 500/500。", CYAN, GREEN_PALE, bsize=9.5)
    card(s, 7.15, 2.98, 5.45, 1.2, "结果 2：跨精度回放",
         "校正后 60/60 vs 未校正 10/60；SPRC 1,148 B = 完整轨迹的 24.76%（R051）。", BLUE, PALE, bsize=9.5)
    card(s, 7.15, 4.34, 5.45, 1.2, "结果 3：BDS 独立确认",
         "独立 seed 40/40、0 false-safe（R059）。", BLUE, PALE, bsize=9.5)
    add_text(s, "局限：单模型单 GPU；无真人盲评（R048 未收集）；无跨硬件验证。",
             7.15, 5.62, 5.45, 0.3, 9.5, True, ORANGE)
    add_text(s, "未来：扩大模型/硬件矩阵；固定长度失败率区间；整数频数合同；盲评与检测实验。",
             7.15, 5.98, 5.45, 0.3, 9.5, True, NAVY)
    add_text(s, "不宣称：通用安全、完全不可检测、跨硬件保证或论文已发表。",
             1.05, 6.7, 11.2, 0.25, 10.5, True, RED, PP_ALIGN.CENTER)
    source(s, "真实图表：deliverables/assets/fig2_main_scale.png（R044，来源 figures/source_data）；结论范围见 RESEARCH_BRIEF.md。")

    prs.save(OUT)
    print("saved:", OUT)
    print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))


if __name__ == "__main__":
    build()
